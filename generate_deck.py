from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    BG_COLOR = RGBColor(15, 23, 42)      # #0f172a Slate dark
    CARD_BG = RGBColor(30, 41, 59)       # #1e293b Card navy
    TEXT_WHITE = RGBColor(255, 255, 255)
    TEXT_MUTED = RGBColor(203, 213, 225) # #cbd5e1
    ACCENT_BLUE = RGBColor(59, 130, 246) # #3b82f6
    ACCENT_GREEN = RGBColor(52, 211, 153)# #34d399

    slides_data = [
        {
            "title": "VoicePulse AI Engine",
            "subtitle": "Enterprise Real-Time Sales QA & Compliance Intelligence Platform",
            "bullets": [
                "Presenter: Pervaiz | Machine Learning & AI Engineer",
                "Core Mission: Safeguarding high-stakes customer calls with sub-second streaming analysis.",
                "Tech Stack: FastAPI, AssemblyAI V3 Universal Streaming, WebSockets, Web Audio API."
            ]
        },
        {
            "title": "The Enterprise Compliance Problem",
            "subtitle": "Why Traditional QA Fails in Fast-Paced Sales Environments",
            "bullets": [
                "Manual Post-Call Bottleneck: Reviews happen after calls conclude, resulting in irreversible regulatory violations.",
                "High Financial Exposure: Absolute guarantees ('100% safe', 'guaranteed returns') lead to massive compliance fines.",
                "Real-Time Intervention Gap: Existing solutions lack low-latency streaming engines capable of whispering real-time alerts to agents."
            ]
        },
        {
            "title": "System Architecture & Data Flow",
            "subtitle": "End-to-End Real-Time Pipeline Design",
            "bullets": [
                "Client Audio Ingestion: Browser Web Audio API captures microphone input, downsampling to 16kHz raw PCM.",
                "Secure Backend Handshake: FastAPI proxy endpoint (`/api/token`) generates short-lived AssemblyAI V3 tokens securely.",
                "Real-Time WebSocket Stream: Bidirectional stream (`wss://`) processes live audio packets sub-second.",
                "Semantic Risk Evaluation: Turn-by-turn POST requests calculate risk scores and trigger instant UI supervisor alerts."
            ]
        },
        {
            "title": "Engineering Challenges & Solutions",
            "subtitle": "Overcoming Complex Integration Roadblocks",
            "bullets": [
                "API Credential Isolation: Routed all token generation requests through a protected backend layer, preventing browser exposure.",
                "Strict Header Sanitization: Implemented rigorous string cleaning (.strip(), quote/newline stripping) to solve 401 auth rejections.",
                "Audio Buffer Optimization: Managed ScriptProcessorNode buffers (`Int16Array`) to ensure zero-latency packet transmission without audio clipping."
            ]
        },
        {
            "title": "Business Impact & Future Roadmap",
            "subtitle": "Measurable Value & Scalability",
            "bullets": [
                "85% Reduction in QA Audit Overhead: Automates compliance checks across thousands of concurrent calls.",
                "Instant Risk Mitigation: Prevents regulatory penalties before a sales agent closes a conversation.",
                "Next-Phase Roadmap: Multi-speaker diarization (Agent vs. Customer segregation) and automated post-call executive PDF reporting."
            ]
        }
    ]

    for data in slides_data:
        slide = prs.slides.add_slide(blank_layout)
        
        # Background fill
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_COLOR
        bg.line.fill.background()

        # Title container
        txBox = slide.shapes.add_textbox(Inches(1.0), Inches(0.8), Inches(11.333), Inches(1.5))
        tf = txBox.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = data["title"]
        p.font.size = Pt(36)
        p.font.bold = True
        p.font.color.rgb = TEXT_WHITE
        p.font.name = "Segoe UI"

        p2 = tf.add_paragraph()
        p2.text = data["subtitle"]
        p2.font.size = Pt(18)
        p2.font.color.rgb = ACCENT_BLUE
        p2.font.name = "Segoe UI"
        p2.space_before = Pt(8)

        # Content Card Background
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(2.3), Inches(11.333), Inches(4.5))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = RGBColor(51, 65, 85)
        card.line.width = Pt(1.5)

        # Content bullets inside card
        c_box = slide.shapes.add_textbox(Inches(1.3), Inches(2.6), Inches(10.7), Inches(3.9))
        c_tf = c_box.text_frame
        c_tf.word_wrap = True

        for i, bullet in enumerate(data["bullets"]):
            bp = c_tf.add_paragraph() if i > 0 else c_tf.paragraphs[0]
            bp.text = f"• {bullet}"
            bp.font.size = Pt(18)
            bp.font.color.rgb = TEXT_MUTED
            bp.font.name = "Segoe UI"
            bp.space_before = Pt(14)

    prs.save("VoicePulse_AI_Presentation.pptx")
    print("Successfully generated VoicePulse_AI_Presentation.pptx!")

if __name__ == "__main__":
    create_deck()